from __future__ import annotations

import hashlib
import os
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.resume_agent.errors import ResumeImportError
from app.resume_agent.json_source import load_json_resume
from app.resume_agent.models import SplitResumeFile


def safe_resolve(root: Path, relative_path: str) -> Path:
    candidate_input = Path(relative_path)
    if candidate_input.is_absolute():
        raise ResumeImportError("INVALID_STORAGE_PATH", "不允许绝对存储路径。")
    root = root.resolve()
    candidate = (root / candidate_input).resolve(strict=False)
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise ResumeImportError("INVALID_STORAGE_PATH", "存储路径越界。") from exc
    current = root
    for part in candidate.relative_to(root).parts:
        current /= part
        if current.exists() and _is_reparse_point(current):
            raise ResumeImportError("REPARSE_POINT_FORBIDDEN", "路径包含重解析点。")
    return candidate


def _is_reparse_point(path: Path) -> bool:
    if path.is_symlink():
        return True
    attributes = getattr(path.stat(), "st_file_attributes", 0)
    return bool(attributes & getattr(os, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def validate_source(path: Path, extension: str, hard_limit: int, json_limit: int = 10 * 1024 * 1024) -> None:
    extension = extension.upper()
    if extension not in {"JSON", "PDF", "DOCX"}:
        raise ResumeImportError("UNSUPPORTED_FILE_TYPE", "系统仅支持JSON、PDF和DOCX简历。")
    if not path.is_file() or path.stat().st_size <= 0:
        raise ResumeImportError("EMPTY_FILE", "简历为空或不存在。")
    if path.stat().st_size > hard_limit:
        raise ResumeImportError("SOURCE_FILE_TOO_LARGE", "简历超过安全上限。")
    if extension == "JSON" and path.stat().st_size > json_limit:
        raise ResumeImportError("JSON_FILE_TOO_LARGE", "JSON简历超过安全上限。")
    try:
        if extension == "JSON":
            load_json_resume(path, min(hard_limit, json_limit))
        elif extension == "PDF":
            with path.open("rb") as stream:
                if stream.read(5) != b"%PDF-":
                    raise ResumeImportError("FILE_STRUCTURE_INVALID", "PDF文件结构无效。")
            PdfReader(str(path))
        else:
            Document(str(path))
    except ResumeImportError:
        raise
    except Exception as exc:
        raise ResumeImportError("FILE_STRUCTURE_INVALID", "简历文件损坏或结构无效。") from exc


def extract_text(path: Path, extension: str) -> str:
    extension = extension.upper()
    if extension == "JSON":
        source_text = load_json_resume(path, 10 * 1024 * 1024).visible_text
    elif extension == "PDF":
        source_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    elif extension == "DOCX":
        document = Document(str(path))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        blocks.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        source_text = "\n".join(blocks)
    else:
        raise ResumeImportError("UNSUPPORTED_FILE_TYPE", "系统仅支持JSON、PDF和DOCX简历。")
    text = normalize_text(source_text)
    if not text:
        raise ResumeImportError("NO_TEXT_CONTENT", "简历没有可读取文本；扫描版PDF请先进行OCR。")
    return text


def split_document(source: Path, extension: str, output_directory: Path) -> list[SplitResumeFile]:
    if extension.upper() != "JSON":
        raise ResumeImportError("UNSUPPORTED_FILE_TYPE", "系统仅支持JSON简历。")
    return [SplitResumeFile(path=str(source), sequence_no=1)]
