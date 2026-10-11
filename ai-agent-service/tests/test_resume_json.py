from __future__ import annotations

import json
from pathlib import Path

import pytest

from app.resume_agent.documents import extract_text, split_document, validate_source
from app.resume_agent.errors import ResumeImportError
from app.resume_agent.json_source import apply_json_identity, build_local_json_profile, build_local_json_resume, clean_candidate_identity, clean_visible_text, load_json_resume, resume_source_from_file_name
from app.resume_agent.models import ParsedResume
from app.resume_agent.optimizer import optimize_if_needed


@pytest.mark.parametrize("extension", ["PDF", "DOCX"])
def test_corrupt_manual_resume_formats_are_rejected(tmp_path: Path, extension: str) -> None:
    source = tmp_path / f"resume.{extension.lower()}"
    source.write_bytes(b"not-used")
    with pytest.raises(ResumeImportError, match="结构"):
        validate_source(source, extension, 10 * 1024 * 1024)


def test_docx_resume_text_can_be_extracted(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("测试候选人")
    document.add_paragraph("五年机械设计经验")
    document.save(source)
    validate_source(source, "DOCX", 10 * 1024 * 1024)
    assert extract_text(source, "DOCX") == "测试候选人\n五年机械设计经验"


def test_json_resume_without_contact_is_valid(tmp_path: Path) -> None:
    source = tmp_path / "51JOB_10001.json"
    source.write_text(json.dumps({
        "talent_id": "10001",
        "source_route": "/Revision/talent/resume/detail",
        "candidate_name": "测试候选人",
        "headline": "电机设计工程师",
        "visible_text": "五年驱动电机结构设计经验，熟练使用UG。",
        "attachment_links": [],
    }, ensure_ascii=False), encoding="utf-8")
    validate_source(source, "JSON", 500 * 1024 * 1024)
    parsed = load_json_resume(source, 10 * 1024 * 1024)
    assert parsed.talent_id == "10001"
    assert parsed.source_route == "/Revision/talent/resume/detail"
    assert resume_source_from_file_name(source.name, parsed.talent_id) == "51JOB"
    assert parsed.candidate_name == "测试候选人"
    assert parsed.headline == "电机设计工程师"
    assert "驱动电机" in parsed.visible_text
    assert "138" not in parsed.visible_text
    assert "@" not in parsed.visible_text
    assert extract_text(source, "JSON") == parsed.visible_text
    parts = split_document(source, "JSON", tmp_path / "split")
    assert len(parts) == 1 and Path(parts[0].path) == source
    path, text, original_size, final_size, optimized = optimize_if_needed(source, "JSON", 20 * 1024 * 1024)
    assert path == source and text == parsed.visible_text
    assert original_size == final_size and optimized is False


def test_json_resume_rejects_missing_visible_text(tmp_path: Path) -> None:
    source = tmp_path / "invalid.json"
    source.write_text('{"talent_id":"10001","source_route":"/Revision/talent/resume/detail","candidate_name":"测试"}', encoding="utf-8")
    with pytest.raises(ResumeImportError, match="visible_text"):
        load_json_resume(source, 10 * 1024 * 1024)


def test_json_resume_rejects_file_name_talent_id_mismatch() -> None:
    with pytest.raises(ResumeImportError, match="人才ID"):
        resume_source_from_file_name("51JOB_99999.json", "10001")


def test_resume_source_from_file_name_accepts_external_job_suffix() -> None:
    assert resume_source_from_file_name("51JOB_10001_172848506.json", "10001") == "51JOB"


def test_json_identity_ignores_contact_patterns_and_uses_file_prefix(tmp_path: Path) -> None:
    source = tmp_path / "51JOB_10001.json"
    source.write_text(json.dumps({
        "talent_id": "10001",
        "source_route": "/Revision/talent/resume/detail",
        "candidate_name": "测试候选人",
        "headline": "电机设计工程师",
        "visible_text": "联系电话13800138000，邮箱test@example.com",
    }, ensure_ascii=False), encoding="utf-8")
    json_source = load_json_resume(source, 10 * 1024 * 1024)
    parsed = ParsedResume(name="AI姓名", phone="13800138000", email="test@example.com", current_title="AI推断职位")
    talent_id, resume_source = apply_json_identity(parsed, json_source, source.name)
    assert (talent_id, resume_source) == ("10001", "51JOB")
    assert parsed.name == "测试候选人"
    assert parsed.current_title == "电机设计工程师"
    assert parsed.phone is None and parsed.email is None


def test_json_local_profile_does_not_require_ai_structuring(tmp_path: Path) -> None:
    source = tmp_path / "51JOB_10001.json"
    source.write_text(json.dumps({
        "talent_id": "10001", "source_route": "/detail", "candidate_name": "测试候选人",
        "headline": "35岁 8年经验 本科", "visible_text": "测试候选人\n35岁 8年经验 本科\n求职意向\n电机设计工程师",
    }, ensure_ascii=False), encoding="utf-8")
    parsed = build_local_json_resume(load_json_resume(source, 10 * 1024 * 1024))
    assert parsed.name == "测试候选人"
    assert parsed.current_title == "电机设计工程师"
    assert parsed.highest_education == "本科" and parsed.total_experience_months == 96
    assert parsed.phone is None and parsed.email is None
    assert parsed.skills == [] and parsed.work_experiences == []


def test_51job_page_shell_and_repeated_watermark_are_removed() -> None:
    company = "测试招聘企业有限公司"
    text = "\n".join([
        "通知", "44", "手机版", "帮助", "企业管理", "人才ID:10001 活跃时间:2026.08.11",
        "扫码分享", "在线简历", "Resume", "测试候选人", "离职（一周内到岗）", "电话沟通",
        "求职意向", "机械工程师", "工作经历", "测试制造公司", f"{company}  {company}  {company}  {company}",
        "声明：以上人才信息仅供测试招聘企业有限公司招聘使用，禁止用于其他任何用途。",
        "操作动态", "暂无动态", "发送",
    ])
    cleaned = clean_visible_text(text)
    assert cleaned == "测试候选人\n离职（一周内到岗）\n求职意向\n机械工程师\n工作经历\n测试制造公司"


def test_company_label_before_masked_name_is_removed_and_name_is_corrected() -> None:
    name, text = clean_candidate_identity(
        "鑫可传动",
        "鑫可传动\n谢先生\n24小时内活跃\n求职意向\n汽车设计工程师\n工作经历\n测试公司",
    )
    assert name == "谢先生"
    assert text.startswith("谢先生\n24小时内活跃\n求职意向")
    assert "鑫可传动" not in text


def test_first_masked_name_is_used_instead_of_source_page_activity_text() -> None:
    name, _ = clean_candidate_identity("李先生 1周内活跃 离职（一周内到岗）", "李先生\n1周内活跃\n离职（一周内到岗）\n求职意向\n质量工程师")
    assert name == "李先生"


def test_database_json_profile_rebuilds_basic_fields_without_ai() -> None:
    profile = build_local_json_profile("李先生 1周内活跃 离职（一周内到岗）", "李先生\n44岁 23年经验 大专 现居·广州\n求职意向\n客户质量工程师")
    assert profile.name == "李先生"
    assert profile.highest_education == "大专"
    assert profile.total_experience_months == 276
    assert profile.current_title == "客户质量工程师"
